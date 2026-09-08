# -*- coding: utf-8 -*-
"""
介護テクノロジー導入支援事業 提案動画 ナレーション台本（Word）+ Vrew用txt
レイアウト: 横向きA4 / 1ページ2スライド / 左サムネイル・右台本 / スライド番号ティール
サムネイル: scratchpad/qa2/スライドN.PNG（PowerPoint書き出し）
"""
import os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn

TEAL=RGBColor(0x02,0x80,0x90); INK=RGBColor(0x1F,0x29,0x37); MUTED=RGBColor(0x64,0x74,0x8B)
QA_DIR=r"C:\Users\MyPC\AppData\Local\Temp\claude\C--Users-MyPC-Desktop-AI-Agents\2b1bb727-cd0b-45eb-8637-bcc966d8caea\scratchpad\qa2"
OUT_DIR=os.path.join("company","secretary","materials")

slides=[
("表紙",
"こんにちは。AIと介護DXの伴走パートナー、「となりにAI」のとーるです。"
"今回は、令和8年度の「介護テクノロジー導入支援事業」についてご紹介します。"
"この補助金を上手に使うと、AIやICTの導入にかかる費用の負担を、大きく減らすことができます。"
"ぜひ最後までご覧ください。"),
("制度の仕組み",
"まず、制度の仕組みからご説明します。"
"この事業は、国が3分の2を負担する基金をもとに、都道府県が実施する補助金です。"
"ポイントは、実施主体が都道府県だということです。"
"そのため、補助率や上限額、募集の時期は、地域ごとに異なります。"
"ご自身の事業所がある都道府県の情報を確認することが、最初の一歩になります。"),
("補助率",
"次に、補助率です。通常の導入では、最大で4分の3が補助されます。"
"さらに、複数の機器をまとめて導入する「パッケージ型」なら、5分の4、つまり8割が補助され、基準額は1,000万円です。"
"ひとつ、注意点があります。"
"令和7年度は、「第三者による業務改善支援」を受けることが、補助の要件になっていました。"
"令和8年度も同じ要件になる可能性がありますので、最新の公募要領をご確認ください。"),
("対象機器・対象経費",
"では、何が補助の対象になるのでしょうか。"
"特に重点的に支援されるのが、介護記録ソフト、見守り機器、インカムの3つです。"
"これらは、業務時間の削減効果が確認されているものです。"
"そして、ここが大きなポイントですが、AIがケアプランの原案づくりを支援するソフトも、対象経費に含まれています。"
"ほかにも、移乗支援や入浴支援の機器などが対象です。"),
("補助上限額の例",
"気になる金額の目安です。"
"たとえば介護記録ソフトの場合、職員数に応じて、100万円から250万円が上限です。"
"タブレットやWi-Fiの整備を含めると、プラス15万円。"
"ケアプランのデータ連携で、プラス5万円の加算もあります。"
"パッケージ型なら、最大1,000万円。"
"さらに、業務改善支援の費用として、45万円から48万円が、別枠で用意されています。"),
("ご提案：段階導入",
"ここからが、私たちからのご提案です。ポイントは、「いきなり全部はやらない」こと。"
"まずステップ1として、記録の音声入力と、AIによる整形から始めます。現場の成功体験づくりです。"
"ステップ2で、見守り機器やインカムを導入し、夜間や連携の負担を減らします。"
"そしてステップ3で、AIによるケアプラン支援へ。"
"ここで得た効果データは、生産性向上推進体制加算の要件にも、そのまま活かせます。"),
("神奈川県の状況",
"神奈川県の状況です。"
"令和8年度の伴走支援の募集は、5月にすでに終了しています。"
"一方で、通常の補助金の申請期間は、まだ公表されていません。"
"つまり今は、公表されたらすぐ動けるように、準備を進めておく時期です。"
"ご相談は、横浜市総合リハビリテーションセンターの介護ロボット相談窓口や、かながわ福祉サービス振興会でも受け付けています。"),
("となりにAIの伴走支援",
"とはいえ、要綱を読み込んで、書類をそろえて、というのは、なかなか大変です。"
"そこで、となりにAIが伴走します。"
"まず、現場の困りごとをお聞きし、現場に合う機器と、使える補助金を一緒に選びます。"
"申請書類や効果測定計画の準備をサポートし、導入後は、職員研修と使い方のフォローまで。"
"「置いていかれる人」を作らない導入を、最後まで支えます。"),
("クロージング",
"最後に、お伝えしたいことがあります。"
"補助金は、機器を安く買うためのものではありません。"
"現場のみなさんが、笑顔で使いこなせるようになるための、時間を買うもの。私たちは、そう考えています。"
"まずは、お気軽にご相談ください。となりにAIの、とーるでした。"
"No Smile, No Life。"),
]

def secs(text):  # 目安: 約330字/分（5.5字/秒）
    s=len(text)/5.5
    return int(round(s/5.0)*5)

def set_run(r,text=None,size=10.5,color=INK,bold=False,font="Meiryo"):
    if text is not None: r.text=text
    r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    r.font.name=font
    r._element.get_or_add_rPr()
    rf=r._element.rPr.get_or_add_rFonts(); rf.set(qn('w:eastAsia'),font)
    return r

doc=Document()
sec=doc.sections[0]
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Mm(297); sec.page_height=Mm(210)
sec.top_margin=Mm(14); sec.bottom_margin=Mm(12); sec.left_margin=Mm(15); sec.right_margin=Mm(15)

total=sum(secs(t) for _,t in slides)
p=doc.add_paragraph(); set_run(p.add_run(),"介護テクノロジー導入支援事業 提案動画　ナレーション台本",16,TEAL,True,"Yu Gothic UI")
p=doc.add_paragraph()
set_run(p.add_run(),f"となりにAI ｜ 全{len(slides)}スライド ｜ 想定尺 約{total//60}分{total%60:02d}秒 ｜ 2026.07　",9.5,MUTED)
set_run(p.add_run(),"※ 秒数は約330字/分の読み上げ目安。Vrew収録時は .txt スクリプトをインポートしてください。",9.5,MUTED)

def block(idx,title,text):
    tbl=doc.add_table(rows=1,cols=2)
    tbl.autofit=False
    for cell,w in zip(tbl.rows[0].cells,(Mm(112),Mm(155))):
        cell.width=w
    left,right=tbl.rows[0].cells
    left.vertical_alignment=WD_ALIGN_VERTICAL.TOP
    lp=left.paragraphs[0]; lp.alignment=WD_ALIGN_PARAGRAPH.LEFT
    img=os.path.join(QA_DIR,f"スライド{idx}.PNG")
    lp.add_run().add_picture(img,width=Mm(108))
    rp=right.paragraphs[0]
    set_run(rp.add_run(),f"スライド {idx}",13,TEAL,True,"Yu Gothic UI")
    set_run(rp.add_run(),f"　{title}",11,INK,True,"Yu Gothic UI")
    set_run(rp.add_run(),f"　（目安 約{secs(text)}秒）",9.5,MUTED)
    bp=right.add_paragraph(); bp.paragraph_format.space_before=Pt(6); bp.paragraph_format.line_spacing=1.35
    set_run(bp.add_run(),text,11,INK)

for i,(title,text) in enumerate(slides,1):
    block(i,title,text)
    if i%2==1 and i<len(slides):
        sp=doc.add_paragraph(); sp.paragraph_format.space_after=Pt(4)
    elif i%2==0 and i<len(slides):
        doc.add_page_break()

os.makedirs(OUT_DIR,exist_ok=True)
docx_path=os.path.join(OUT_DIR,"介護テクノロジー導入支援事業_ナレーション台本.docx")
doc.save(docx_path)

txt_path=os.path.join(OUT_DIR,"介護テクノロジー導入支援事業_Vrewスクリプト.txt")
with open(txt_path,"w",encoding="utf-8") as f:
    f.write("\n\n".join(t for _,t in slides)+"\n")
print("SAVED"); print(docx_path); print(txt_path); print("total_sec",total)
